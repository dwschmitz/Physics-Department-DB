import io
import unittest

import openpyxl

import importer
import physdb

HEADER = ["First Name", "Last Name", "UCID", "CNET", "Email", "Citizenship", "Matriculation Year (7th year clock)",
          "Matriculation Quarter", "Transfer", "Transfer Year", "Transfer Quarter"]


def xlsx(rows, sheet_title="Incoming Student Data "):
    wb = openpyxl.Workbook()
    wb.active.title = sheet_title
    for r in rows:
        wb.active.append(r)
    buf = io.BytesIO(); wb.save(buf)
    return buf.getvalue()


def good(n=1, **over):
    row = dict(zip(HEADER, ["Ada%d" % n, "Lovelace%d" % n, 10000000 + n, "ada%d" % n, "ada%d@u.edu" % n, "Citizen", 2026,
                            "Autumn", 0, "-", "-"]))
    row.update(over)
    return [row[h] for h in HEADER]


class ImporterTests(unittest.TestCase):
    def setUp(self):
        self.conn = physdb.connect(":memory:")

    def run_rows(self, *rows, header=HEADER):
        return importer.analyse(self.conn, xlsx([header, *rows]), "f.xlsx")

    def test_the_example_file_imports_cleanly(self):
        r = importer.analyse(self.conn, open("examples/Incoming_Student_Data_Test.xlsx", "rb").read(), "x.xlsx")
        self.assertEqual((len(r["new"]), len(r["duplicates"]), len(r["problems"]), r["ignored"]), (7, 0, 0, []))
        by = {x["v"]["last_name"]: x["v"] for x in r["new"]}
        self.assertEqual(by["Feynman"]["citizenship"], "international")
        self.assertEqual((by["Feynman"]["transfer_year"], by["Feynman"]["transfer_term"]), (None, None))
        pauli = by["Pauli"]
        self.assertEqual((pauli["matriculation_year"], pauli["matriculation_term"], pauli["transfer_year"], pauli["transfer_term"]),
                         (2024, "Autumn", 2026, "Autumn"))                  # trailing space in "Autumn " is trimmed
        self.assertEqual(by["Davis"]["university_id"], "90876543")
        self.assertEqual(r["new"][0]["row"], 2)                              # row numbers are the sheet's own

    def test_good_row_values(self):
        r = self.run_rows(good(1, CNET="ADA1", Citizenship=" permanent resident "), good(2, **{"Matriculation Quarter": "winter "}))
        self.assertEqual(len(r["new"]), 2)
        v = r["new"][0]["v"]
        self.assertEqual((v["cnet"], v["citizenship"]), ("ada1", "permanent_resident"))
        self.assertEqual(r["new"][1]["v"]["matriculation_term"], "Winter")

    def test_ucid_must_be_eight_digits(self):
        for bad in (9876543, 123456789, "ABC12345", "", None, 1234.5):
            r = self.run_rows(good(1, UCID=bad))
            self.assertEqual((len(r["new"]), len(r["problems"])), (0, 1), bad)
            self.assertIn("UCID", r["problems"][0]["messages"][0])
        r = self.run_rows(good(1, UCID=12345678.0), good(2, UCID="23456789"))      # Excel floats and text are fine
        self.assertEqual([x["v"]["university_id"] for x in r["new"]], ["12345678", "23456789"])

    def test_row_problems(self):
        cases = {"First Name": "", "Last Name": " ", "Email": "no-at-sign", "Citizenship": "Martian",
                 "Matriculation Year (7th year clock)": "twenty", "Matriculation Quarter": "Fall", "CNET": "bad cnet!"}
        for col, val in cases.items():
            r = self.run_rows(good(1, **{col: val}))
            self.assertEqual((len(r["new"]), len(r["problems"])), (0, 1), col)
        r = self.run_rows(good(1, Citizenship=""))
        self.assertIn("Citizenship is empty", r["problems"][0]["messages"][0])
        r = self.run_rows(good(1, **{"Matriculation Year (7th year clock)": 1899}))
        self.assertEqual(len(r["problems"]), 1)
        r = self.run_rows(good(1, Email="", UCID=1))                              # every problem on the row is listed
        self.assertEqual(len(r["problems"][0]["messages"]), 2)

    def test_transfer_rules(self):
        ok = self.run_rows(good(1, Transfer=1, **{"Transfer Year": 2026, "Transfer Quarter": "Autumn "}),
                           good(2, Transfer="yes", **{"Transfer Year": "2025", "Transfer Quarter": "spring"}),
                           good(3, Transfer=0, **{"Transfer Year": "-", "Transfer Quarter": "-"}),
                           good(4, Transfer=None), good(5, Transfer=0, **{"Transfer Year": 2026, "Transfer Quarter": "Autumn"}))
        self.assertEqual(len(ok["new"]), 5)
        flags = [(x["v"]["transfer_year"], x["v"]["transfer_term"]) for x in ok["new"]]
        self.assertEqual(flags, [(2026, "Autumn"), (2025, "Spring"), (None, None), (None, None), (None, None)])
        for kw in (dict(Transfer=1), dict(Transfer=1, **{"Transfer Year": 2026}), dict(Transfer=1, **{"Transfer Quarter": "Autumn"}),
                   dict(Transfer=1, **{"Transfer Year": "-", "Transfer Quarter": "-"}), dict(Transfer="maybe")):
            r = self.run_rows(good(1, **kw))
            self.assertEqual((len(r["new"]), len(r["problems"])), (0, 1), kw)

    def test_duplicates_are_skipped_not_updated(self):
        physdb.add_student(self.conn, university_id="10000001", first_name="Old", last_name="Name", email="old@u.edu",
                           matriculation_year=2020, research_area="kept", notes=None)
        r = self.run_rows(good(1, **{"First Name": "Changed"}), good(2))
        self.assertEqual((len(r["new"]), len(r["duplicates"]), len(r["problems"])), (1, 1, 0))
        d = r["duplicates"][0]
        self.assertEqual((d["row"], d["who"], d["existing"]["last_name"]), (2, "Lovelace1, Changed", "Name"))
        self.assertEqual(self.conn.execute("SELECT first_name, research_area FROM students WHERE university_id = '10000001'").fetchone()[:],
                         ("Old", "kept"))                                       # analysing never writes
        # a duplicate is skipped even if the rest of its row is wrong
        r = self.run_rows(good(1, Citizenship="Martian"))
        self.assertEqual((len(r["duplicates"]), len(r["problems"])), (1, 0))

    def test_email_and_cnet_clashes_are_problems(self):
        physdb.add_student(self.conn, university_id="10000001", first_name="A", last_name="B", email="Taken@u.edu",
                           matriculation_year=2020, research_area=None, notes=None, cnet="taken")
        r = self.run_rows(good(2, Email="taken@U.edu"), good(3, CNET="TAKEN"), good(4))
        self.assertEqual((len(r["new"]), len(r["problems"])), (1, 2))
        self.assertIn("different student", r["problems"][0]["messages"][0])
        self.assertIn("(B, A)", r["problems"][1]["messages"][0])

    def test_repeats_inside_the_file(self):
        r = self.run_rows(good(1), good(2, UCID=10000001), good(3, Email="ADA1@u.edu"), good(4, CNET="ada1"), good(5))
        self.assertEqual([x["row"] for x in r["new"]], [2, 6])
        self.assertEqual([x["row"] for x in r["problems"]], [3, 4, 5])
        self.assertIn("first on row 2", r["problems"][0]["messages"][0])
        r = self.run_rows(good(1), good(1))                                       # same student twice, both new
        self.assertEqual((len(r["new"]), len(r["problems"])), (1, 1))

    def test_headings_are_matched_loosely_and_extras_ignored(self):
        header = ["first name", "LAST NAME", "University ID", "cnet", "E-mail", "citizenship", "Matriculation year", "matriculation term",
                  "Favourite colour", "Notes", "Research Area"]
        r = importer.analyse(self.conn, xlsx([header, ["A", "B", 12345678, "ab", "ab@u.edu", "Citizen", 2025, "Winter", "red", "n", "optics"]]), "f.xlsx")
        self.assertEqual((len(r["new"]), r["ignored"]), (1, ["Favourite colour"]))
        v = r["new"][0]["v"]
        self.assertEqual((v["notes"], v["research_area"], v["transfer_year"]), ("n", "optics", None))   # no Transfer columns at all

    def test_missing_columns_and_bad_files(self):
        with self.assertRaises(importer.ImportFileError) as cm:
            importer.analyse(self.conn, xlsx([["First Name", "Last Name"], ["A", "B"]]), "f.xlsx")
        for name in ("UCID", "Email", "Citizenship", "Matriculation Year", "Matriculation Quarter"):
            self.assertIn(name, str(cm.exception))
        for data, name, word in ((b"x", "f.xlsx", "valid Excel"), (b"", "f.csv", "no data"), (b"a", "f.pdf", "upload an Excel"),
                                 (b"a", "f.xls", "old-style"), (xlsx([]), "f.xlsx", "no data")):
            with self.assertRaises(importer.ImportFileError, msg=name) as cm:
                importer.analyse(self.conn, data, name)
            self.assertIn(word, str(cm.exception))

    def test_blank_rows_header_position_and_only_first_sheet(self):
        data = xlsx([[None] * 3, HEADER, good(1), [None] * 11, good(2), [" ", ""]])
        r = importer.analyse(self.conn, data, "f.xlsx")
        self.assertEqual(([x["row"] for x in r["new"]], r["header_row"], r["total"]), ([3, 5], 2, 2))
        wb = openpyxl.Workbook(); wb.active.append(HEADER); wb.active.append(good(1))
        wb.create_sheet("Second").append(["junk"]); buf = io.BytesIO(); wb.save(buf)
        self.assertEqual(len(importer.analyse(self.conn, buf.getvalue(), "f.xlsx")["new"]), 1)

    def test_csv(self):
        text = ",".join(HEADER) + "\nAda,Lovelace,12345678,ada,ada@u.edu,citizen,2026,Autumn,0,-,-\n" \
               "José,O'Brien,23456789,,jose@u.edu,International,2026,Autumn,1,2026,Winter\n"
        for raw in (text.encode("utf-8"), ("﻿" + text).encode("utf-8"), text.encode("cp1252")):
            r = importer.analyse(self.conn, raw, "f.csv")
            self.assertEqual((len(r["new"]), len(r["problems"])), (2, 0))
            self.assertEqual(r["new"][1]["v"]["last_name"], "O'Brien")
        self.assertEqual(r["new"][1]["v"]["first_name"], "José")
        t = importer.csv_template()
        self.assertEqual(importer.analyse(self.conn, (t + "Ada,L,12345678,,a@u.edu,Citizen,2026,Autumn,0,,\n").encode("utf-8"), "t.csv")["new"][0]["v"]["last_name"], "L")

    def test_limits(self):
        many = xlsx([HEADER] + [good(i) for i in range(1, importer.MAX_ROWS + 2)])
        with self.assertRaises(importer.ImportFileError) as cm:
            importer.analyse(self.conn, many, "f.xlsx")
        self.assertIn("more than", str(cm.exception))


if __name__ == "__main__":
    unittest.main()
